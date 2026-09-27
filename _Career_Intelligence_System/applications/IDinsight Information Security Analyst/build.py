#!/usr/bin/env python3
"""Information Security Analyst, IDinsight. Close 15 Oct 2026. Kenya office.
Do not claim Security+, SASE/CASB tenant admin, DPO, or ISO 27001 lead.
CV 1-2 pages. No photo, no DOB. Canonical PATH titles.
Run: python3 "applications/IDinsight Information Security Analyst/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT = (
    "Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "Information Security Analyst  |  EDR, MDM, SIEM-class ops, SETA, Kenya DPA  |  INGO",
"metrics": [
    ("3 yrs", "IT security ops"),
    ("EDR", "MDM  SIEM"),
    ("80+", "SETA trained"),
    ("Kenya", "no visa"),
],
"show_referees": True,
"compact_education": True,
"summary": ("Information security analyst with 3 years of IT security operations at PATH, an INGO: "
    "hands-on administration of EDR, MDM-class endpoint management and SIEM-class monitoring, plus "
    "firewall and identity controls in the same family as SASE and CASB. I investigate alerts, run "
    "SETA, write policies people actually use, assess risk on systems and data flows, and support "
    "Kenya Data Protection Act and GDPR-aligned practice. Kenyan national. BSc Applied Computing, "
    "cybersecurity and digital forensics."),
"skills": [
    ("Security stack", "EDR and endpoint protection; MDM-class fleet (imaging, group policy, coverage across a distributed estate); SIEM-class logs, detections and alerting (Grafana, system and security logs); SASE/CASB-class control via identity, VPN and firewall policy (FortiGate, Meraki MX); MFA, IAM, patching"),
    ("IR, risk, GRC", "Incident triage, investigation, containment, recovery and post-incident write-up; information and data security risk assessments; vulnerability reporting; ISO/IEC 27001 and NIST CSF control thinking; risk register and assurance evidence; partner due diligence packs"),
    ("SETA, policy, privacy", "Security education, training and awareness; onboarding security training; phishing-awareness and higher-risk team sessions; policies, standards and guidelines that get used; Kenya Data Protection Act 2019, GDPR principles (mapping, retention, subject rights); AI governance and acceptable-use controls"),
    ("Systems and people", "Internal help desk (Freshdesk, Jira Service Management); Microsoft 365 and Azure; Active Directory / Entra ID; written and spoken English; cross-cultural, multi-site teams"),
],
"experience": [
    {"org": "PATH", "title": "Health Informatics / HMIS Officer", "dates": "Mar 2025 to Present",
     "loc": "Kenya (INGO, multi-site)", "bullets": [
        "Own day-to-day security operations on a distributed fleet: EDR coverage, SIEM-class log and alert review, MDM-class device compliance, and identity/MFA policy; investigate alerts and write what changes after the incident.",
        "Run information and data security risk assessments on systems, vendors and data flows, including health and programme data, and turn findings into actions teams can take inside their timelines.",
        "Support Kenya Data Protection Act and GDPR-aligned practice (access, retention, evidence) and AI acceptable-use guidance so responsible use is the default; keep an evidence pack for due-diligence questions.",
    ]},
    {"org": "PATH", "title": "ICT Associate", "dates": "Jan 2024 to Mar 2025",
     "loc": "Kenya (INGO, multi-site)", "bullets": [
        "Administered endpoint protection/EDR, patching, MFA and access control; tuned firewall and VPN policy (the control layer next to SASE/CASB); maintained coverage and investigated security events under SLA.",
        "Designed and ran SETA: onboarding security training, awareness material, and sessions for staff handling sensitive data; 80+ people trained so the policy changed behaviour, not just the slide deck.",
        "Wrote SOPs and security guidelines that were adopted on the queue; contributed to the control set, audits and the help desk so tooling stayed useful rather than merely deployed.",
    ]},
    {"org": "PATH", "title": "ICT Intern", "dates": "Apr 2023 to Dec 2023",
     "loc": "Kenya", "bullets": [
        "Supported logging, tickets, firewall and VLAN work with senior engineers, and wrote up procedures used in later incident handling.",
    ]},
],
"certs": "AZ-900  |  Google IT Support  |  IBM Technical Support  |  CCNA  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "Re: Information Security Analyst, IDinsight (Kenya)",
    "",
    "I am applying for the Information Security Analyst role, Kenya office. I am a Kenyan national. I do not need visa sponsorship. I can start as soon as possible and I can commit two years.",
    "For three years at PATH I have been the person who makes the security stack do work: EDR, MDM-class fleet management, SIEM-class monitoring, identity and firewall policy. I tune detections, keep coverage on a distributed estate, investigate alerts, contain, recover, and write the post-incident note. I design SETA that staff actually sit through. I assess risk on systems, vendors and data flows and I hand project teams something they can do this month. I keep policies and evidence so the next due-diligence request is faster. I know Kenya's Data Protection Act and GDPR principles from both study and a health-data environment. I will follow ISO/IEC 27001 and NIST CSF as the language of the control set. I will sit Security+ as the first cert you fund.",
    "I have not been the named SASE or CASB tenant owner. I already do that job on the identity, VPN and firewall layer, and I get productive on a new console the way I did on every other one: with the alert, with the policy, and with the user. I still take help-desk tickets. That is how I know whether the control is usable.",
    "Put the stack in front of me and I will run it.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]


def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "Information Security Analyst")
    br.add_contact(doc, CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)


if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
