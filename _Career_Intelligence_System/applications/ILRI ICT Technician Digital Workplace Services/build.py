#!/usr/bin/env python3
"""ICT Technician - Digital Workplace Services, ILRI Kenya.
REF: ICT-CS/2308-09/2026. Deadline 28 Sep 2026. Canonical PATH titles.
Do not claim Intune tenant admin, MD-102, or Copilot rollout lead.
Run: python3 "applications/ILRI ICT Technician Digital Workplace Services/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT_NAIROBI = (
    "Nairobi / Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "ICT Technician - Digital Workplace Services  |  Microsoft 365, endpoints, user adoption  |  Nairobi",
"metrics": [
    ("3 yrs", "end-user computing"),
    ("120+", "users supported"),
    ("~98%", "availability"),
    ("80+", "staff trained"),
],
"show_referees": True,
"compact_education": True,
"referees": [
    "Patrick Ondieki, ICT Officer, PATH  |  pondieki@path.org",
    "Marvin Ngosa, Systems Administrator, PATH  |  mngosa@path.org",
    "Gadaffi Ochieng, Procurement Officer, Ampath  |  ochienggadaffi@gmail.com",
],
"summary": ("ICT Technician with about three years of ICT customer support and end-user computing at "
    "PATH, a global health organisation with multi-site Kenyan operations. First- and second-line digital "
    "workplace services for 120+ users at roughly 98% availability: Microsoft 365 (Teams, SharePoint "
    "Online, OneDrive, Outlook), Windows endpoints, Microsoft Entra ID and MFA, onboarding and user "
    "training. I clear incidents and service requests against agreed service levels, document the fix, "
    "and trace recurring causes. Kenyan national, BSc Applied Computing, AZ-900."),
"skills": [
    ("Digital workplace", "Microsoft 365 - Teams, SharePoint Online, OneDrive, Outlook, Admin Center; collaboration technologies, digital productivity, secure use"),
    ("Identity and access", "Microsoft Entra ID / Azure AD, Active Directory, authentication, MFA, provisioning, onboarding and offboarding"),
    ("Endpoint management", "Windows 10/11 laptops, desktops, mobile devices; imaging and provisioning, software deployment, patch management, device compliance, endpoint security, asset recovery, Intune / Configuration Manager equivalent"),
    ("ITSM and service levels", "Freshdesk, Zoho Desk, Jira Service Management; incidents and service requests, escalation, SLA, knowledge articles and documentation"),
    ("Adoption and enablement", "Orientation sessions, digital workplace awareness, productivity demonstrations, Microsoft Copilot / AI-enabled workplace services, 80+ staff trained; Teams meetings and room kit for hybrid workshops and institutional events"),
],
"experience": [
    {"org": "PATH", "title": "Health Informatics / HMIS Officer", "dates": "Mar 2025 to Present",
     "loc": "Kenya (Hybrid)", "bullets": [
        "First-line and second-line Microsoft 365 and digital productivity support across Teams, SharePoint Online, OneDrive and Outlook: restore the service, escalate complex identity or infrastructure issues with diagnostics attached.",
        "Onboard new users end to end - accounts, licences, MFA, endpoint devices - run orientation on collaboration and secure use, and keep knowledge articles and user guides current.",
        "Technical support for hybrid meetings, workshops and institutional events (Teams, room AV) with multidisciplinary teams.",
    ]},
    {"org": "PATH", "title": "ICT Associate", "dates": "Jan 2024 to Mar 2025",
     "loc": "Kenya (enterprise, multi-site)", "bullets": [
        "Owned the digital workplace queue for 120+ users on Freshdesk against agreed service levels: response time down about 25%, availability near 98% across a multi-site environment.",
        "Owned the Windows endpoint lifecycle: provisioning and imaging, software deployment, patch management, compliance and security checks, preventive maintenance, asset recovery and accurate ICT asset records, using Microsoft imaging, group policy and deployment tooling.",
        "Administered Microsoft Entra ID / Active Directory and Microsoft 365 licences, permissions and MFA; ran user orientation and productivity demonstrations, and analysed recurring incidents to remove causes.",
    ]},
    {"org": "PATH", "title": "ICT Intern", "dates": "Apr 2023 to Dec 2023",
     "loc": "Kenya", "bullets": [
        "Windows device support, imaging, Freshdesk incidents and end-user documentation under ITIL practice and SLA.",
    ]},
],
"certs": "AZ-900  |  Google IT Support  |  IBM Technical Support  |  CCNA  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Head of People and Culture,",
    "",
    "Re: ICT Technician - Digital Workplace Services, REF: ICT-CS/2308-09/2026",
    "",
    "I am a Kenyan national applying for the ICT Technician - Digital Workplace Services position at ILRI Nairobi. For about three years at PATH, a global health organisation running multi-site operations in Kenya, this has been my daily work: first-line and second-line digital workplace support for more than 120 users across Microsoft 365 (Teams, SharePoint Online, OneDrive, Outlook), Windows endpoints, Microsoft Entra ID and MFA, and onboarding. I hold availability near 98%, resolve incidents and service requests against agreed service levels on Freshdesk, escalate complex issues with the diagnostic evidence attached, and write the knowledge article afterwards.",
    "On endpoints I handle the full lifecycle: provisioning and imaging, software deployment, patch management, security and compliance checks, preventive maintenance, asset recovery, and accurate ICT asset records. I deliver this with Microsoft imaging, group policy and deployment tooling, which is the equivalent endpoint management experience your Terms of Reference call for, and it is the same lifecycle Intune and Configuration Manager automate, so I move into either console quickly.",
    "What I would bring beyond closing tickets is the part your Terms of Reference emphasises: adoption and improvement. I have trained more than 80 staff through orientation sessions and productivity demonstrations, written the guides and SOPs that stop repeat questions, and worked the causes behind recurring incidents instead of the symptoms. I support Teams meetings and workshop kit, and I would carry Microsoft Copilot and other approved AI-enabled workplace services the same way: orientation, clear guidance, and a feedback loop back to the team that owns the platform.",
    "I hold a BSc in Applied Computing from KCA University and Microsoft Azure Fundamentals (AZ-900). I am available for Nairobi on this three-year HG12 contract and would welcome the chance to discuss how I can support ILRI's digital workplace.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]


def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "ICT Technician - Digital Workplace Services")
    br.add_contact(doc, CONTACT_NAIROBI)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)


if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT_NAIROBI
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
