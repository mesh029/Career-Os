#!/usr/bin/env python3
"""ICT Trainer — St. Kizito Vocational Centre Utawala. Deadline 16 Sep 2026.
Facts per 00_Master_Resume.md. Do not claim TSC or a teaching diploma.
Run: python3 "applications/St Kizito Vocational Centre ICT Trainer/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT_NAIROBI = (
    "Nairobi / Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "ICT Trainer  |  Digital skills, practical labs and equipment  |  Vocational / youth",
"metrics": [
    ("80+", "people trained"),
    ("15+", "sessions"),
    ("120+", "users supported"),
    ("ICT", "labs and kit"),
],
"show_referees": False,
"compact_education": True,
"summary": ("ICT Trainer with about three years at PATH teaching digital skills to people who are not "
    "IT: 80+ staff across 15+ sessions, plus the lab work that makes a class possible. I deliver "
    "practical training, sit with learners on project work, and keep the equipment and learning "
    "resources running so the session does not die on a dead printer. Windows, Microsoft 365, "
    "networks, SOPs, how-to guides. BSc Applied Computing (KCA University, 2023). Ready for "
    "Utawala, Nairobi, full-time."),
"skills": [
    ("ICT training", "Digital skills delivery, practical sessions, how-to guides, patient English with new learners, 80+ people trained"),
    ("Project work", "Support learners on real tasks: accounts, files, networks, simple applications, documentation"),
    ("Equipment and labs", "Windows 10/11, printers, imaging, LAN/Wi-Fi, Microsoft 365, asset inventory, keep the room usable"),
    ("Learning resources", "SOPs, job aids, knowledge base, simple exercises, checklists after class"),
    ("Support behind the class", "1st/2nd-line help desk, Freshdesk, MFA, identity, backups, vendor follow-up"),
    ("Education", "BSc Applied Computing, Cybersecurity and Digital Forensics; AZ-900, Google IT Support"),
],
"experience": [
    {"org": "PATH", "title": "ICT Trainer", "dates": "Mar 2025 to Present",
     "loc": "Kenya (Hybrid)  ·  ready to relocate to Nairobi", "bullets": [
        "Deliver digital-skills training to non-technical staff: Microsoft 365, secure use, data tools. 80+ people, 15+ sessions. Practical, not a lecture.",
        "Sit with learners on project work until the task works: accounts, files, reports, the thing they have to submit.",
        "Maintain ICT equipment and learning resources so the session starts: Windows, printers, network, access, the how-to on the wall.",
    ]},
    {"org": "PATH", "title": "ICT Associate", "dates": "Jan 2024 to Mar 2025",
     "loc": "Kenya (multi-site)", "bullets": [
        "Trained users while running 1st- and 2nd-line support for 120+ people. Response about 25% faster once the how-to existed.",
        "Kept labs and endpoints usable: imaging, printers, Microsoft 365, Active Directory, basic network faults.",
        "Wrote knowledge-base articles and SOPs that doubled as teaching materials.",
    ]},
    {"org": "PATH", "title": "ICT Intern", "dates": "Apr 2023 to Dec 2023",
     "loc": "Kenya", "bullets": [
        "Supported equipment, imaging, printers, and tickets. Learned to set the room up before people arrive.",
    ]},
],
"certs": "AZ-900  |  Google IT Support  |  IBM Technical Support  |  CCNA  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the ICT Trainer post at St. Kizito Vocational Centre, Utawala. I want to teach ICT as a usable skill, not as a slide deck. For about three years at PATH I have trained more than 80 people in digital tools, sat with them on the practical work, and kept the equipment running so the session did not fail on a cable.",
    "That is this job. Deliver ICT and digital skills training. Support practical sessions and project work. Maintain ICT equipment and learning resources. I hold a BSc in Applied Computing from KCA University (November 2023), plus Azure Fundamentals and Google IT Support. I do not hold a TSC number. I teach by doing, and I can relocate to Nairobi for this post.",
    "I respect the Centre's work with youth who need a trade, not a performance. I can start as soon as you need someone in the lab.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]


def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "ICT Trainer")
    br.add_contact(doc, CONTACT_NAIROBI)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)


if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], CONTACT_NAIROBI)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], CONTACT_NAIROBI, cover_title)
    print("Generated:", HERE)
