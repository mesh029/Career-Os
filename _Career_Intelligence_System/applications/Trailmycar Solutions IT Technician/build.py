#!/usr/bin/env python3
"""IT Technician, Trailmycar Solutions. Deadline 23 Sep 2026.
Email hr@trailmycar.com subject IT Technician.
Do not claim a telematics/GPS console. CRM = Salesforce tickets. Degree is KCA not JKUAT.
Run: python3 "applications/Trailmycar Solutions IT Technician/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT = (
    "Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "IT Technician  |  Helpdesk, dashboards, device health, client support  |  CRM",
"metrics": [
    ("3 yrs", "IT support"),
    ("120+", "users"),
    ("~25%", "faster response"),
    ("Daily", "ticket summaries"),
],
"show_referees": True,
"compact_education": True,
"summary": ("IT Technician with 3 years of IT support and helpdesk at PATH: inbound tickets, hardware, "
    "software and network issues, user accounts and company devices, dashboards for device health, "
    "technical reports, backups and IT security. I log the case, troubleshoot, resolve or escalate, "
    "train the user, and send daily and weekly summaries of what was open, what was offline, and what "
    "was done. Salesforce CRM tickets. BSc Applied Computing, KCA University."),
"skills": [
    ("Helpdesk and client support", "Inbound tickets and calls, Freshdesk, Zoho Desk, Jira Service Management, Salesforce CRM tickets, log-troubleshoot-resolve-or-escalate, SLA, recurring-issue notes, daily and weekly case summaries"),
    ("Dashboards and reports", "Device-health and performance dashboards (Grafana, Power BI, Tableau, Meraki), flag offline or underperforming devices, trip-style movement and offline logs, technical reports for users and clients"),
    ("Internal IT", "Hardware, software and network support, user accounts and access, company device setup, antivirus, backups, IT security controls, IT equipment, licences and credentials inventory, third-party vendor coordination"),
    ("Training", "Staff and user training on systems, dashboards and tools, job aids, 80+ people trained"),
],
"experience": [
    {"org": "PATH", "title": "Health Informatics / HMIS Officer", "dates": "Mar 2025 to Present",
     "loc": "Kenya (multi-site)", "bullets": [
        "Provide IT support to internal staff and programme users: hardware, software, network, accounts and access, with tickets logged and closed against SLA.",
        "Monitor dashboards for device and service health, flag offline or underperforming endpoints, and generate technical reports and summaries for the people who asked.",
        "Support system updates, antivirus, backups and IT security controls, and keep equipment, licence and credential records accurate.",
    ]},
    {"org": "PATH", "title": "ICT Associate", "dates": "Jan 2024 to Mar 2025",
     "loc": "Kenya (multi-site)", "bullets": [
        "Owned the inbound support queue for 120+ users on Freshdesk: log, troubleshoot, resolve or escalate field work, cutting response time about 25%.",
        "Set up user accounts, configured company devices, managed access including Salesforce CRM tickets (login, access, forms and reports), and coordinated vendors on escalations.",
        "Watched network and endpoint dashboards for offline devices, trained staff on the tools, and prepared daily and weekly summaries of cases, device health and action taken.",
    ]},
    {"org": "PATH", "title": "ICT Intern", "dates": "Apr 2023 to Dec 2023",
     "loc": "Kenya", "bullets": [
        "Logged Freshdesk tickets, supported desktops, imaging and network issues, and documented recurring faults for the knowledge base.",
    ]},
],
"certs": "Google IT Support  |  IBM Technical Support  |  CCNA  |  AZ-900  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "Re: IT Technician",
    "",
    "I am applying for the IT Technician role. For three years at PATH this has been the job: inbound technical cases, internal IT support, dashboards for device health, reports, and a clear log of what was done. I take the call or the ticket, I log it, I fix it or I escalate for field intervention, and I send the summary.",
    "I have supported 120+ users: hardware, software, network, user accounts, company devices, access, antivirus, backups and security. I watch dashboards, flag offline or underperforming devices, and produce the technical report the user asked for. I have handled Salesforce CRM tickets. I train people on the tools so the same question does not come back. Ticket response improved about 25%.",
    "I have not administered a vehicle-tracking console. I have done the work around one: client-facing diagnostics, dashboard monitoring, offline-device flags, reports, and a helpdesk that stays calm on the phone. I pick up TMC's platform the same way I picked up every other console: with the user on the line and the ticket open.",
    "I hold a BSc in Applied Computing from KCA University. I am available as soon as possible.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]


def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "IT Technician")
    br.add_contact(doc, CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)


if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
