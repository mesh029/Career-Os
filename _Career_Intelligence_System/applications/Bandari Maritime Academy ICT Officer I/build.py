#!/usr/bin/env python3
"""ICT Officer I — Bandari Maritime Academy (BMA/26/CEO/ICT1/17).
Facts per 00_Master_Resume.md. Canonical PATH titles. Do not claim PSC grade or MCSE.
Run: python3 "applications/Bandari Maritime Academy ICT Officer I/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT_NAIROBI = (
    "Nairobi / Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "ICT Officer  |  Hardware, software, networks and user support  |  AZ-900",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("80+", "staff trained"),
    ("AZ-900", "Microsoft cert"),
],
"show_referees": True,
"compact_education": True,
"referees": [
    "Patrick Ondieki — ICT Officer, PATH — pondieki@path.org",
    "Marvin Ngosa — Systems Administrator, PATH — mngosa@path.org",
    "Gadaffi Ochieng — Procurement Officer, Ampath — ochienggadaffi@gmail.com",
],
"summary": ("ICT Officer with around three years at PATH installing, configuring and maintaining "
    "computer hardware, software and networks for 120+ users at roughly 98% uptime. Provides user "
    "support and training, writes technical and user documentation, and keeps Windows, Microsoft 365 "
    "and Active Directory running. Protects information and the network (MFA, least privilege, patching). "
    "Writes and tests programmes to specification (Python, SQL, Next.js). Holds Microsoft Azure Fundamentals "
    "(AZ-900) and completed CCNA coursework. Ready to relocate to Nairobi."),
"skills": [
    ("Hardware and software", "Install, configure and maintain Windows 10/11 PCs, imaging, printers, Microsoft 365, day-to-day ICT platforms"),
    ("Networks and security", "LAN/WAN, DNS, DHCP, VPN, Cisco Meraki, MikroTik, CCNA coursework; MFA, Active Directory, least privilege, endpoint protection, firewalls"),
    ("User support and training", "1st/2nd-line support, Freshdesk/Zoho Desk, ITIL/SLA, 80+ staff trained, clear English"),
    ("Documentation and systems", "Technical, systems and user documentation, SOPs, knowledge base, implement systems, ICT resource and asset registers"),
    ("Programming to spec", "Write and test Python, SQL and Next.js applications and APIs to instructions; REST integrations"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid)","bullets":[
        "Operate and maintain ICT platforms alongside systems work for 120+ users, keeping critical services above 98% uptime.",
        "Implement and support computer systems, write technical and user documentation, and train staff on secure, correct use.",
        "Protect information and the network: access control, MFA, backups, encryption, and vulnerability follow-up.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Installed, configured and maintained computer hardware, Windows 10/11 software, and ICT infrastructure for 120+ users. Ticket response improved about 25%.",
        "Provided user support and training; wrote knowledge-base articles and SOPs so staff could keep using the systems after the fix.",
        "Managed day-to-day platforms: Active Directory, Microsoft 365, Cisco Meraki networks (LAN/WAN, DNS, DHCP, VPN, Wi-Fi).",
        "Supported ICT resource control and procurement inputs: asset inventory, licence awareness, vendor and ISP specifications when equipment or connectivity was due.",
    ]},
    {"org":"Independent / Freelance","title":"Digital Systems & IT Support","dates":"2023 to Present","loc":"Remote (project-based)","bullets":[
        "Wrote and tested programmes and APIs to specification (Next.js, TypeScript, Python, SQL), including a platform with 40+ REST APIs, then documented how to run it.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported hardware, LAN/WAN, imaging, Freshdesk tickets, and process documentation under ITIL and SLA.",
    ]},
],
"certs": "AZ-900  |  CCNA coursework (Cybrary)  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Human Resource,",
    "",
    "I am applying for ICT Officer I at Bandari Maritime Academy, reference BMA/26/CEO/ICT1/17. For about three years at PATH I have done the work this grade describes in a comparable ICT role: install and configure hardware and software, keep the network and day-to-day platforms running, support and train users, write the documentation, and protect information. I supported 120+ users at around 98% uptime. I trained 80+ staff. I hold a BSc in Applied Computing from KCA University and Microsoft Azure Fundamentals (AZ-900). I completed CCNA coursework.",
    "I have not held Assistant ICT Officer II in the public service. PATH ICT Associate and ICT support is the comparable period. I do not hold MCSE or the Cisco CCNA exam. AZ-900 is my Microsoft certification. The attached CCNA document is Cybrary course completion.",
    "I write and test programmes to specification where the Academy needs it (Python, SQL, Next.js), and I can gather requirements and specifications for ICT equipment with vendors. I implement ICT workplans and standards. I have not written an institutional maritime ICT strategy, and I will not pretend I have.",
    "I will relocate to Nairobi for a permanent, on-site Officer I post. Thank you.",
    "",
    "Yours faithfully,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "ICT Officer I")
    br.add_contact(doc, CONTACT_NAIROBI)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT_NAIROBI
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
