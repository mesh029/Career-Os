#!/usr/bin/env python3
"""IT Microsoft Support Specialist — BruntWork Seattle healthcare client.
20h/week, Seattle business hours. Facts per 00_Master_Resume.md.
Do not claim HIPAA certified, MS-900, or CompTIA A+.
Run: python3 "applications/BruntWork IT Microsoft Support Specialist/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "IT Microsoft Support Specialist  |  Microsoft 365, Entra ID, help desk  |  Healthcare remote",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("M365", "Teams / SharePoint"),
    ("20h", "Seattle hours"),
],
"show_referees": False,
"compact_education": True,
"summary": ("IT Microsoft Support Specialist with about three years at PATH, a healthcare organisation, "
    "supporting 120+ users at roughly 98% uptime. First-line Microsoft 365: Outlook, Teams, SharePoint, "
    "OneDrive, Office apps, plus Windows, printers, and basic connectivity. I own the ticket: password "
    "resets, MFA, onboarding and offboarding, licences and permissions, then I write the how-to so the "
    "CEO is not answering the same question twice. Entra ID / Azure AD, Microsoft 365 Admin Center, "
    "Freshdesk. Remote support. Clear English with non-technical staff. Available 20 hours a week inside "
    "Seattle business hours."),
"skills": [
    ("Microsoft 365", "Admin Center, Outlook, Teams, SharePoint, OneDrive, Office apps, licences, onboarding and offboarding"),
    ("Identity / access", "Entra ID / Azure AD, Active Directory, MFA, password resets, role-based permissions, least privilege"),
    ("Help desk", "Ticketing, prioritise, escalate to vendors or management, remote desktop / TeamViewer-style sessions, SLA"),
    ("Windows / hardware", "Windows 10/11, desktops, peripherals, printers, basic network connectivity, asset inventory, warranty and repair"),
    ("Security / privacy", "Phishing awareness, unauthorised-access watch, MFA, healthcare data hygiene, SOPs and how-to guides"),
    ("Delivery", "Report to leadership, independent remote work, train non-technical users, AZ-900"),
],
"experience": [
    {"org": "PATH", "title": "Health Informatics / HMIS Officer", "dates": "Mar 2025 to Present",
     "loc": "Kenya (Hybrid / Remote)  ·  healthcare", "bullets": [
        "Primary Microsoft 365 support for healthcare staff: Outlook, Teams, SharePoint, OneDrive, accounts, MFA, and access. Keep identity and mail working at over 98% uptime.",
        "Onboarding and offboarding: provision accounts, licences, and permissions by role. Leavers lose access the same day.",
        "Write how-to guides and SOPs for recurring questions so non-technical users, and leadership, are not stuck on the same ticket.",
    ]},
    {"org": "PATH", "title": "ICT Associate", "dates": "Jan 2024 to Mar 2025",
     "loc": "Kenya (multi-site healthcare)", "bullets": [
        "First-line help desk for 120+ users: Freshdesk tickets, password resets, MFA, Windows, printers, and basic connectivity. Response about 25% faster under SLA.",
        "Administered Microsoft 365 and Active Directory / Entra ID. Remote sessions when the user was not in the room. Escalated to vendors or senior IT when the break was outside the runbook.",
        "Kept a live hardware and licence inventory. Trained 80+ staff on secure use so phishing and leftover access did not become the next incident.",
    ]},
    {"org": "PATH", "title": "ICT Intern", "dates": "Apr 2023 to Dec 2023",
     "loc": "Kenya", "bullets": [
        "Supported Windows desktops, printers, imaging, and Freshdesk tickets. Learned to write what was already tried before handing a case up.",
    ]},
    {"org": "Independent", "title": "Microsoft 365 setup (project-based)", "dates": "2023 to Present",
     "loc": "Remote", "bullets": [
        "Set up Microsoft 365 and Google Workspace for small businesses: accounts, onboarding, and the short guide the next admin could follow.",
    ]},
],
"certs": "AZ-900  |  Google IT Support  |  IBM Technical Support  |  CCNA  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the IT Microsoft Support Specialist role supporting your Seattle healthcare client. For about three years at PATH I have been first-line Microsoft 365 support for more than 120 healthcare users: Outlook, Teams, SharePoint, OneDrive, Windows, printers, MFA, and the ticket that arrives as panic and leaves as a how-to. Uptime stayed around 98%. I already explain this to non-technical staff, and I already report up, not sideways.",
    "I own onboarding and offboarding: accounts, licences, permissions, same-day leavers. I watch phishing and leftover access. I keep the asset list and the SOP. I escalate to vendors or management when the break is outside the runbook. I hold Microsoft Azure Fundamentals (AZ-900), Google IT Support, and IBM Technical Support. I have not sat MS-900 or CompTIA A+. I have not been a US HIPAA officer. I already protect health information through access control, MFA, and training, and I will follow your privacy rules.",
    "I am based in Kenya, fully remote, and available Monday to Friday, 20 hours inside Seattle business hours (a working block for me is 9:00 AM to 1:00 PM Pacific, which is 7:00 to 11:00 PM my time). I can start immediately.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]


def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "IT Microsoft Support Specialist")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)


if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
