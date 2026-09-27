#!/usr/bin/env python3
"""IT Operations Specialist — BruntWork (SimplePractice / Stripe audit seat).
Facts per 00_Master_Resume.md. Do not claim SimplePractice, Stripe, or HIPAA training.
Run: python3 "applications/BruntWork IT Operations Specialist/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "IT Operations Specialist  |  Platform audit, configuration & documentation  |  Healthcare-aware",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("40+", "REST APIs"),
    ("SOPs", "docs / KB"),
],
"show_referees": False,
"compact_education": True,
"summary": ("IT operations specialist with around three years at PATH reviewing and configuring "
    "healthcare and workplace platforms for 120+ users at roughly 98% uptime. Audits settings, access, "
    "and data flow, then corrects configuration and writes the SOP so another team can keep it. Integrates "
    "systems (REST APIs, EMR–lab–pharmacy, workflow automation) and configures Paystack so systems can take "
    "payments, including M-Pesa STK push. Also reviews websites and SaaS setups (Microsoft 365, Zoho, "
    "Google Workspace, WordPress). Protects health information through access control "
    "and least privilege. Remote-ready for US Central Time hours."),
"skills": [
    ("Platform audit", "Review settings, access, forms, workflows, and paid-feature use; correct configuration; keep affiliated sites consistent"),
    ("Integrations", "REST APIs, EMR–lab–pharmacy and national data-warehouse links, workflow automation, data accuracy"),
    ("Payments and SaaS", "Paystack API, M-Pesa STK push, multi-channel collection; Microsoft 365, Zoho, Google Workspace, WordPress / CMS, technical site review"),
    ("Documentation", "SOPs, job aids, knowledge base, process notes the next person can follow"),
    ("Healthcare privacy", "Access control, encryption, least privilege, staff training; HIPAA-aware, ready to sign BAA"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid / Remote)","bullets":[
        "Audit healthcare platforms (KenyaEMR/TaifaCare, DHIS2): settings, access, data flow, and risk. Correct configuration and keep critical systems above 98% uptime.",
        "Review integrations between clinical and reporting systems (EMR–laboratory–pharmacy, national data-warehouse APIs) and tighten workflows so data stays accurate.",
        "Write SOPs, job aids, and technical notes so programme and facility teams can maintain the set-up after the change.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Configured Microsoft 365, Active Directory, and Freshdesk / Zoho Desk for 120+ users across sites: accounts, permissions, and the same standard in each place.",
        "Turned repeat tickets into knowledge-base articles and onboarding steps so the process, not the person, carried the work.",
        "Worked independently and remotely with non-technical staff; isolated identity vs endpoint vs network, then documented the fix.",
    ]},
    {"org":"Independent / Freelance","title":"Digital Systems & IT Support","dates":"2023 to Present","loc":"Remote (project-based)","bullets":[
        "Reviewed and set up small-business platforms: Microsoft 365, Zoho, Google Workspace, and websites (WordPress / CMS). Checked what was misconfigured, unused, or inconsistent.",
        "Configured Paystack APIs so more than one system could collect payments from different channels, including M-Pesa STK push. Built sites and workflows (including a Next.js platform with 40+ REST APIs) and left documentation the client could keep.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported Freshdesk tickets, imaging, and process documentation under ITIL and SLA.",
    ]},
],
"certs": "AZ-900  |  CCNA  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the Information Technology Operations Specialist role. For about three years at PATH I have reviewed and configured healthcare and workplace systems for 120+ users: audit the set-up, fix what is wrong, document it so the team can keep it. That is the job you described, on a different named stack.",
    "I have not administered SimplePractice or Stripe. I have configured Paystack payment APIs, including M-Pesa STK push, so more than one system can take payments from different channels. That is the same job as reviewing Stripe: keys, payment methods, test vs live, and the path money takes. I have not run a social scheduler as the owner. I already audit configuration, websites, and integrations, and I can live in SimplePractice, Stripe, and your scheduler from week one.",
    "I work in healthcare IT. I protect health information through access control, encryption, and least privilege. I am not HIPAA-course certified. I will sign and follow Business Associate and privacy agreements.",
    "I am based in Kenya, fully remote, and available Monday to Friday, 9:00 AM to 6:00 PM US Central. I hold a BSc in Applied Computing plus AZ-900, CCNA, Google IT Support, and IBM Technical Support.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "Information Technology Operations Specialist")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
