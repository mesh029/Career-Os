#!/usr/bin/env python3
"""Technical Customer Support Analyst — Cigna Nairobi.
Facts per 00_Master_Resume.md. Do not claim Citrix, SCCM, ServiceNow, or Avaya.
Run: python3 "applications/Cigna Technical Customer Support Analyst/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT_NAIROBI = (
    "Nairobi / Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "Technical Customer Support Analyst  |  Service Desk, IMAC, Active Directory  |  Healthcare IT",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("~25%", "faster incidents"),
    ("SLA", "Freshdesk / ITIL"),
],
"show_referees": False,
"compact_education": True,
"summary": ("Service Desk IT support professional with around three years at PATH supporting 120+ "
    "onsite and remote users in a healthcare setting at roughly 98% uptime. Owns incidents and service "
    "requests in Freshdesk and Zoho Desk under ITIL and SLA, does IMAC work on Windows 10/11 PCs "
    "(install, image, configure, hardware), and manages user access in Active Directory and Microsoft 365 "
    "(including mailbox / Exchange Online). Patches endpoints and troubleshoots network faults, then "
    "escalates outages with a clean handover. Hybrid-ready for Nairobi office and home rota."),
"skills": [
    ("Service Desk", "Incidents and service requests, Freshdesk, Zoho Desk, Jira Service Management, ITIL workflows, SLA, prioritise, escalate, knowledge base"),
    ("IMAC and Windows 11", "Install, move, add, change; Windows 10/11 OS and PC hardware; PXE imaging; printers; onsite and remote support"),
    ("Access and Exchange", "Active Directory, group policy, MFA, least privilege, Microsoft 365 / Office, Outlook, Exchange Online mailbox access"),
    ("Networking and patching", "LAN/WAN, DNS, DHCP, VPN, Cisco Meraki, endpoint patching, hardware faults"),
    ("Healthcare users", "Clinical and programme staff, clear English, hybrid delivery"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid / Remote)","bullets":[
        "Provide technical support to onsite and remote healthcare users: accounts, Windows devices, access, and service restoration, keeping critical systems above 98% uptime.",
        "Manage user access in Active Directory and Microsoft 365 (mailbox / Exchange Online, Teams, MFA) so joiners work and leavers lose access the same day.",
        "Escalate outages with reproduction notes and the layer already ruled out (identity, endpoint, network, server). Write the SOP when the same incident will return.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Owned the Service Desk queue for 120+ users in Freshdesk and Zoho Desk: log incidents and service requests, hit SLA, prioritise, escalate. Ticket response improved about 25%.",
        "IMAC on end-user devices: install and configure Windows 10/11 (OS and PC hardware), Microsoft 365 / Office, and peripherals; image desktops (PXE) for standard builds.",
        "Administered Active Directory (provisioning, rights, group policy) and Exchange Online / Outlook mailbox access; supported MFA and endpoint protection.",
        "Patched endpoints and diagnosed network faults (DNS, DHCP, VPN, Cisco Meraki). First call when connectivity dropped; ISP when the fault was upstream.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported LAN/WAN, PC hardware, printers, VoIP, PXE imaging, and Freshdesk tickets under ITIL and SLA.",
    ]},
],
"certs": "AZ-900  |  CCNA  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the Technical Customer Support Analyst role in Nairobi. For about three years at PATH I have been the Service Desk: onsite and remote technical support for 120+ healthcare users, incidents and service requests in Freshdesk and Zoho Desk under SLA, IMAC on Windows 10/11 PCs, Active Directory and Microsoft 365 access (including Exchange Online mailboxes), networking and patching. Uptime stayed around 98%. Ticket response improved about 25%.",
    "Cigna is a health insurer. I already support people who cannot wait for a vague ticket. I install and build Windows devices, unlock and provision AD, keep the queue clean, and escalate outages to the next layer with facts. I hold a BSc in Applied Computing and Microsoft Azure Fundamentals (AZ-900). I work ITIL in production. I do not hold ITIL Foundation.",
    "I have not administered Citrix, SCCM or ServiceNow. I already do the jobs those tools exist for: remote support, imaging and patching, and a call management system. I can live in the Cigna consoles from week one. I use 3CX for telephony, not Avaya.",
    "I will relocate for the 66% office / 34% home pattern. I can work Saturday 8 AM to 1 PM every three weeks and Monday to Friday 4 PM to midnight every six weeks, from home, if that rota is still current.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "Technical Customer Support Analyst")
    br.add_contact(doc, CONTACT_NAIROBI)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT_NAIROBI
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
