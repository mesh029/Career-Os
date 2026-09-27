#!/usr/bin/env python3
"""Tailored CV + Cover Letter for CMMB Kenya IT Specialist (volunteer).
Facts per 00_Master_Resume.md. Canonical PATH titles. No ECHIS.
Run: python3 "applications/CMMB Kenya IT Specialist Volunteer/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "IT Specialist  |  Health NGO first-line support  |  Microsoft 365, facilities and digital health",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("80+", "staff trained"),
    ("PATH", "health facilities"),
],
"show_referees": False,
"compact_education": True,
"summary": ("IT specialist with around three years at PATH supporting 120+ users across Kenya health "
    "facilities at roughly 98% uptime. First-line and escalated support: Microsoft 365, Active Directory "
    "and Entra ID, Windows devices, backups, and security awareness. Trains non-technical clinical and "
    "programme staff, keeps asset and helpdesk records accurate, and supports health information systems "
    "(KenyaEMR / DHIS2) and reporting dashboards. Kenya-based and ready to volunteer with a country office "
    "or partner facility."),
"skills": [
    ("Service desk", "First-line incidents, Freshdesk, Zoho Desk, ITIL, SLA, knowledge base, clear English with non-technical users"),
    ("Identity and Microsoft 365", "Entra ID / Active Directory, MFA, least privilege, Teams, SharePoint, Outlook, onboarding and offboarding"),
    ("Devices and field sites", "Windows 10/11, imaging, patching, endpoint protection, LAN/WAN, Cisco Meraki, remote support"),
    ("Health systems", "KenyaEMR / TaifaCare, DHIS2, SQL validation, Power BI / Grafana, staff training on digital tools"),
    ("Security and assets", "Awareness training, access control, encryption, backups and recovery, IT asset inventory, licences"),
    ("NGO delivery", "Multi-site Kenya, SOPs, vendor coordination, confidentiality, volunteer and intern enablement"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid)","bullets":[
        "Support health facilities with first-line and escalated IT: EMR/HMIS availability, Windows and Linux servers, backups, encryption, and recovery, above 98% uptime.",
        "Administer Microsoft 365 and Entra ID / Active Directory (accounts, MFA, least privilege) and train facility and programme staff on secure use of digital tools.",
        "Build SQL and dashboard reporting (Power BI / Grafana) that cut manual reporting time about 40%, and keep SOPs and operational logs current.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Delivered first- and second-line support to 120+ users across sites, about 25% faster ticket response under Freshdesk / Zoho SLA.",
        "Configured Windows endpoints, Microsoft 365, and Active Directory; monitored Cisco Meraki; supported VoIP and device setup for new staff.",
        "Kept the asset inventory and knowledge base accurate; ran security awareness and MFA so joiners were not the weak point.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported LAN/WAN, imaging, Freshdesk tickets, and documentation for non-technical users at health sites.",
    ]},
],
"certs": "CCNA  |  AZ-900  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Volunteer Programme Team,",
    "",
    "I am applying for the IT Specialist volunteer role listed under CMMB Kenya. I live in Kenya and I am asking to serve as a local volunteer with your Kenya team, not as an international volunteer based in the United States.",
    "For the past three years at PATH I have supported 120+ users across health facilities, keeping services at over 98% uptime. The work is first-line IT in a health setting: Microsoft 365 and identity, device setup, backups, helpdesk discipline, and training clinicians and programme staff who are not IT people. I also support KenyaEMR and DHIS2 and reporting tools. I have not used ECHIS.",
    "CMMB’s values (collaboration, excellence, accountability, respect) match how I already work: finish the ticket, write the SOP, protect health information, and leave the next person a runbook. I can commit six to twelve months if the placement fits, and I can relocate within Kenya for the office or facility you need.",
    "Thank you for considering a Kenya-based IT specialist who already works in health.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "IT Specialist (Volunteer)")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
