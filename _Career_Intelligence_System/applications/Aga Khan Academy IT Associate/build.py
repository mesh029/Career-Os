#!/usr/bin/env python3
"""Tailored CV + Cover Letter for Aga Khan Academy Nairobi IT Associate (4954BR).
Facts per 00_Master_Resume.md. Canonical PATH titles only.
Run: python3 "applications/Aga Khan Academy IT Associate/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT_NAIROBI = (
    "Nairobi / Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "IT Associate  |  First-line support, infrastructure and cybersecurity  |  Schools and multi-site",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("80+", "staff trained"),
    ("SLA", "helpdesk / ITIL"),
],
"show_referees": False,
"compact_education": True,
"summary": ("IT support professional with around three years at PATH delivering first-line and escalated "
    "end-user support for 120+ people across several sites at roughly 98% uptime. Keeps the helpdesk "
    "accurate, meets SLAs, and explains technical issues in clear English. Supports networks, Microsoft 365, "
    "telephony, devices, CCTV-related systems, identity and access, backups, and security awareness. "
    "Trains non-technical users, tracks assets and licences, and coordinates vendors. Ready to work on the "
    "Aga Khan Academy Nairobi campus."),
"skills": [
    ("Service desk", "First-line incidents and requests, Freshdesk, Zoho Desk, ITIL, SLA, knowledge base, customer service"),
    ("Infrastructure", "LAN/WAN, Cisco Meraki, Wi-Fi, Microsoft 365, Azure (AZ-900), Windows devices, VoIP/3CX, AV and teaching-room kit, CCTV and access control"),
    ("Cybersecurity", "Identity and MFA, least privilege, endpoint protection, vulnerability reporting, awareness training, incident support, data protection"),
    ("Systems and data", "Approved applications, licensing, Power BI and Grafana dashboards, SQL reporting, documentation"),
    ("Assets and vendors", "IT asset inventory, software licences, procurement support, vendor and ISP coordination, project workstreams"),
    ("Continuity", "Backups, recovery, patching, encryption, business continuity routines"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid)","bullets":[
        "Provide first-line and escalated IT support alongside systems work for 120+ users, keeping critical services above 98% uptime.",
        "Administer Microsoft 365 and Active Directory / Entra ID (accounts, MFA, least privilege); support Windows endpoints, backups, encryption, and recovery.",
        "Support cybersecurity hygiene: access control, vulnerability reporting, incident handling, and staff awareness. Train users on secure use of digital tools.",
        "Keep operational logs, SOPs, and helpdesk records current; support dashboards (Power BI / Grafana) used for reporting.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Resolved first- and second-line incidents and service requests for 120+ users under Freshdesk / Zoho Desk ITIL and SLA, improving ticket response about 25%.",
        "Configured and monitored Cisco Meraki networks; diagnosed LAN/WAN, Wi-Fi, DNS, DHCP, and VPN faults; supported VoIP/3CX and classroom-style AV setups.",
        "Ran onboarding and offboarding: devices, Active Directory, Microsoft 365, licences. Kept the IT asset inventory and knowledge base accurate.",
        "Supported endpoint protection, patching, MFA, and vendor or ISP coordination when the next layer was outside the team.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported LAN/WAN, printers, VoIP, desktop imaging, Freshdesk tickets, and IT documentation for non-technical users.",
    ]},
],
"certs": "CCNA  |  AZ-900  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis  |  Cyberise and NRF cybersecurity hackathons",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the IT Associate role at the Aga Khan Academy, Nairobi (4954BR). For the past three years at PATH I have provided first-line and escalated IT support for more than 120 users across several sites, keeping services at over 98% uptime. That work is helpdesk discipline: log the ticket, meet the SLA, explain the fix in plain English, and write the SOP so the same issue does not bounce back.",
    "I support the stack this role lists: networks (Cisco Meraki, LAN/WAN, Wi-Fi), Microsoft 365 and identity, Windows devices, telephony (3CX), teaching-room AV, and security and surveillance systems. I help keep cybersecurity controls real on the ground: MFA, least privilege, endpoint protection, vulnerability follow-up, awareness training, and incident support. I also track IT assets and software licences, coordinate vendors, and run backups and recovery so a bad day has a restore path.",
    "I have not administered a school IB learning platform. I have trained 80+ non-technical staff on digital tools and I treat student and staff data as confidential. I hold a BSc in Applied Computing (Cybersecurity and Digital Forensics), plus CCNA, Microsoft Azure Fundamentals (AZ-900), Google IT Support, and IBM Technical Support. I am based in Kisumu and ready to relocate to Nairobi for a campus role.",
    "Thank you for the chance to keep Academy staff and students working.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "IT Associate")
    br.add_contact(doc, CONTACT_NAIROBI)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT_NAIROBI
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
