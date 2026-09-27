#!/usr/bin/env python3
"""IT Administrator — Greenspoon Kenya. Facts per 00_Master_Resume.md.
Run: python3 "applications/Greenspoon Kenya IT Administrator/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

CONTACT_NAIROBI = (
    "Nairobi / Kisumu, Kenya  |  +254 741 174 779  |  aririmeshack@gmail.com  |  "
    "linkedin.com/in/meshack-ariri  |  meshreallycodes.online"
)

cv = {
"headline": "IT Administrator  |  Helpdesk, assets, licences, Microsoft 365 and network  |  Nairobi-ready",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("~25%", "faster tickets"),
    ("Day-1", "joiners / leavers"),
],
"show_referees": False,
"compact_education": True,
"summary": ("IT administrator with around three years at PATH owning the work this role describes: "
    "helpdesk end to end, live asset records, software access and licences, Microsoft 365 and Active Directory, "
    "office network and WiFi, MFA, endpoint protection, and backups. Supports 120+ users at about 98% uptime, "
    "with about 25% faster ticket response. Writes tickets and procedures in clear English. Ready to work in Nairobi."),
"skills": [
    ("Helpdesk", "End-to-end queue, Freshdesk, Zoho Desk, ITIL, SLA, same-day joiners and leavers, knowledge base"),
    ("Assets and licences", "Device register, condition and assignment, licence and subscription tracking, unused-seat flags"),
    ("Accounts and access", "Microsoft 365, Entra ID / Active Directory, MFA, least privilege, join / change / leave"),
    ("Network and WiFi", "TCP/IP, DNS, DHCP, VPN, Cisco Meraki, MikroTik, guest vs staff, ISP escalation"),
    ("Security", "Endpoint protection, patching, firewalls, backups, access control, Data Protection Act 2019 awareness"),
    ("Written English", "Tickets, SOPs, asset notes as the record of what happened"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid)","bullets":[
        "Own escalated IT for 120+ users: accounts, Windows devices, backups, encryption, and recovery, keeping services above 98% uptime.",
        "Keep access aligned to who still works here: Microsoft 365 and Entra ID / Active Directory, MFA, least privilege, joiner and leaver on the day.",
        "Maintain asset awareness, SOPs, and helpdesk notes so a spot check of devices and access has a paper trail.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Ran the support queue for 120+ users under Freshdesk / Zoho SLA: log, resolve, escalate with notes; about 25% faster response.",
        "Configured Microsoft 365, Active Directory, group policy, MFA, and endpoint protection; set up joiners and recovered kit from leavers.",
        "Kept the IT asset inventory current; tracked software access; monitored Cisco Meraki / MikroTik WiFi and LAN; called the ISP when the fault was upstream.",
        "Ran backups, patching, and firewall hygiene so the same control gap did not sit open.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported LAN/WAN, routers, switches, WiFi, imaging, Freshdesk tickets, and IT documentation.",
    ]},
],
"certs": "CCNA  |  AZ-900  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Greenspoon hiring team,",
    "",
    "I am applying for the IT Administrator role. For about three years at PATH I have owned the systems other teams depend on: laptops and assets, licences and access, Microsoft 365, the network and WiFi, MFA, backups, and a helpdesk I run to the end, not to a handoff. About 120+ users, around 98% uptime, about 25% faster ticket response.",
    "That maps to how you measure the job. I keep a live asset record. I flag unused access and licences. I set up joiners and collect kit from leavers on the day. I segregate guest and staff traffic and I am the first call when WiFi dies. I write tickets and procedures in English that a non-IT colleague can follow. I treat Kenyan Data Protection Act 2019 as access, encryption, and no leftover accounts, not as a poster.",
    "I have not run a Mac fleet. Windows is my daily estate. I have set up Google Workspace for small businesses; PATH is Microsoft 365. I will relocate to Nairobi.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "IT Administrator")
    br.add_contact(doc, CONTACT_NAIROBI)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    original = br.CONTACT
    br.CONTACT = CONTACT_NAIROBI
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    br.CONTACT = original
    print("Generated:", HERE)
