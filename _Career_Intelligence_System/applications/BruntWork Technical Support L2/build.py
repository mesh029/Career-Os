#!/usr/bin/env python3
"""Tailored CV + Cover Letter (DOCX + PDF) for Technical Support Agent L2
(BruntWork healthcare MSP client). Facts per 00_Master_Resume.md.
Run from repo: python3 "applications/BruntWork Technical Support L2/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "Technical Support Agent Level 2  |  Healthcare IT, Patching & Network  |  3CX · HIPAA-aware",
"metrics": [
    ("120+", "users supported"),
    ("~98%", "uptime"),
    ("80+", "staff trained"),
    ("3CX", "VoIP / telephony"),
],
"show_referees": False,
"compact_education": True,
"summary": ("Level 2 IT support professional with around three years at PATH supporting 120+ users across "
    "healthcare sites at roughly 98% uptime. Owns escalated troubleshooting: Windows and Linux patching, "
    "Active Directory and Microsoft 365, and network diagnostics (LAN/WAN, VPN, DNS/DHCP, Cisco Meraki). "
    "Protects health information through access control, encryption, and staff training, and will sign "
    "Business Associate agreements. Writes SOPs, supports VoIP/3CX, and automates repeat work with SQL and Python. "
    "Remote-ready for Eastern Time hours."),
"skills": [
    ("L2 / Escalation", "RCA, remote support, ITIL/SLA, Freshdesk, Zoho Desk, ticket quality, knowledge base"),
    ("Patching & Config", "Windows 10/11 and Server, Linux, endpoint protection, MFA, backups, disaster recovery"),
    ("Identity / M365", "Active Directory / Entra ID, group policy, least privilege, Exchange Online, Teams, SharePoint"),
    ("Networking", "LAN/WAN, VLANs, VPN, DNS/DHCP, Cisco Meraki, MikroTik, firewalls, Wi-Fi; CCNA-backed"),
    ("Healthcare IT", "Clinical systems, health-data protection, encryption, HIPAA-aware, ready to sign BAA"),
    ("MSP Tooling", "SOPs, Meraki monitoring, VoIP/3CX, SQL/Python automation, ready to ramp Enable/N-able"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid / Remote)","bullets":[
        "Own escalated technical issues for healthcare sites: EMR/HMIS availability, Windows and Linux server patching, configuration, backups, encryption, and recovery, sustaining over 98% uptime.",
        "Troubleshoot network and access problems (LAN/WAN, VPN, DNS, DHCP, identity) and explain fixes clearly to non-technical clinical and program staff.",
        "Build SQL and Python validation/automation that cut manual reporting and processing time by about 40%, the same habit I would apply to repeat MSP fixes.",
    ]},
    {"org":"PATH","title":"ICT Associate","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya","bullets":[
        "Delivered 1st- and 2nd-line support to 120+ users across multiple healthcare sites, reaching around 98% uptime and about 25% faster ticket response under SLA.",
        "Configured and patched Windows 10/11 endpoints, Microsoft 365, and Active Directory (provisioning, group policy, least privilege, MFA).",
        "Configured and monitored Cisco Meraki network devices; diagnosed LAN/WAN, VPN, Wi-Fi, DNS, and DHCP faults instead of guessing at the symptom.",
        "Supported VoIP/3CX telephony alongside Freshdesk/Zoho Desk ITIL workflows, knowledge-base articles, and standardized procedures.",
    ]},
    {"org":"PATH","title":"ICT Intern","dates":"Apr 2023 to Dec 2023","loc":"Homa Bay, Kenya","bullets":[
        "Supported LAN/WAN, VoIP, desktop imaging, and Freshdesk tickets under ITIL/SLA; contributed to IT documentation.",
    ]},
],
"certs": "AZ-900  |  CCNA  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis  |  Cyberise & NRF cybersecurity hackathons",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the Technical Support Agent Level 2 role supporting your healthcare managed services clients. For the past three years at PATH I have provided Level 2-style IT support in a healthcare environment: escalated troubleshooting, Windows and Linux patching, system configuration, and network diagnostics for more than 120 users, keeping services at over 98% uptime.",
    "That work already looks like this job. I isolate whether an outage is identity, endpoint, network, or server, then I fix it, document it, and turn repeat fixes into SOPs so the same ticket does not bounce back. I administer Active Directory and Microsoft 365 with MFA and least privilege, monitor Cisco Meraki networks, support VoIP/3CX, and use SQL and Python automation to remove repetitive work. I am used to explaining technical issues to clinicians and program staff who are not IT people.",
    "I want to be direct on two points. My production experience is in-house healthcare IT across multiple sites, not a named US MSP running Enable or N-able yet. The outcomes are the same jobs those RMM consoles exist to do, and I can ramp the console quickly. Second, I have not been a US HIPAA officer, but I already protect health information through access control, encryption, backups, and staff training, and I will sign and follow Business Associate and data security agreements.",
    "I am based in Kenya, fully remote-ready, and available Monday through Friday, 9:00 AM to 1:00 PM Eastern. I hold a BSc in Applied Computing (Cybersecurity and Digital Forensics) plus CCNA, Microsoft Azure Fundamentals (AZ-900), Google IT Support, and IBM Technical Support. Thank you for the chance to take escalations off the business owner and keep healthcare clients stable.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "Technical Support Agent Level 2")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
