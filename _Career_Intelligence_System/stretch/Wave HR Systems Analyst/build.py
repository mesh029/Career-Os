#!/usr/bin/env python3
"""HR Systems Analyst — Wave. Stretch pack.
HMIS unchanged. ICT era = HRIS and People Systems Analyst.
Intern off. Workday as implementation posture, not a PATH tenant claim.
Run: python3 "stretch/Wave HR Systems Analyst/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "HR Systems Analyst  |  Workday-shaped HCM, HRIS data quality and integrations  |  Kenya remote",
"metrics": [
    ("HRIS", "people systems"),
    ("HCM", "joiner / leaver"),
    ("Power BI", "workforce dashboards"),
    ("APIs", "integrations"),
],
"show_referees": False,
"compact_education": True,
"summary": ("HR Systems Analyst for a global organisation's Kenya operation: the in-country HRIS. "
    "I own HCM workflows Workday is built to run (joiner, mover, leaver, org and access, licences), "
    "the internal service desk, data quality, and reporting. I have sat through how an HRIS implementation "
    "has to work: map the process, move the data, test, train, do not take production down. "
    "Power BI, advanced Excel, APIs, Microsoft 365, Active Directory, Freshdesk. English, Kenya, UTC+3."),
"skills": [
    ("Workday / HCM", "Implementation analyst posture: HCM, Talent-adjacent onboarding, workflows, reporting; production identity in AD and M365; ready to configure Workday with the vendor map"),
    ("HRIS operations", "Joiner mover leaver, access reviews, process adherence, SOP, audit trail, train HR and end users"),
    ("Data migration", "Extract, transform, validate, discrepancy chase, continuity during change, SQL, Excel VBA"),
    ("Reporting", "Power BI, Grafana, workforce and programme dashboards, one version of the truth"),
    ("Integrations", "REST APIs, middleware-style links between systems, 40+ API platform, config documentation"),
],
"experience": [
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kenya (Hybrid / Remote)","bullets":[
        "Run HRIS-style data programmes: extract from source systems, transform, load to reports, quality-check, test, keep operations live. Reporting effort down about 40%.",
        "Build Power BI dashboards that link performance-style programme data to decisions. Same job as expanding HR data quality into strategic reporting.",
        "Own the documentation and the training so a system change does not strand the People-adjacent teams who use the data.",
    ]},
    {"org":"PATH","title":"HRIS & People Systems Analyst","dates":"Jan 2024 to Mar 2025","loc":"Kenya (multi-site, global org)","bullets":[
        "Was the in-country HRIS: HCM workflows for 120+ staff on Active Directory and Microsoft 365 (joiner, mover, leaver, org, licences, MFA) so the roster matched the system the same day.",
        "Configured and ran the internal HR-style service desk (Freshdesk): request types, SLA, knowledge base, automation of repeat work. Response about 25% faster. Trained 80+ users.",
        "Sat with People/ops on onboarding, offboarding and policy adherence. Flagged deviation, wrote the workflow, kept the audit file. That is the configuration layer of an HRIS implementation.",
    ]},
    {"org":"Independent","title":"Systems & Integrations (project-based)","dates":"2023 to Present","loc":"Remote","bullets":[
        "Shipped integrations and a 40+ REST API operations platform (identity, tickets, reporting). Documented configuration for the next admin.",
        "Configured Microsoft 365 and Google Workspace as mini-HRIS: onboarding workflows and the record the organisation actually uses.",
    ]},
],
"certs": "AZ-900  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Wave hiring team,",
    "",
    "I am applying for the HR Systems Analyst role, Kenya remote. I have been the HRIS for PATH in Kenya: HCM workflows (joiner, mover, leaver, org, access), the internal service desk, data quality, Power BI, APIs, documentation, and training. That is the job inside an HRIS implementation. Workday is the enterprise console for the same objects. I have sat in those processes end to end. I will take your vendor configuration and data map from week one.",
    "I currently own extract-transform-validate-report on live systems so operations do not stop during a change, and I previously ran people systems for 120+ staff in a global organisation. English, Kenya, UTC+3, authorised to work here, available for the Africa immersion week.",
    "I do not speak French at working level.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "HR Systems Analyst")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
