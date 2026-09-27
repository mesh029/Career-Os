#!/usr/bin/env python3
"""Tier 3 IT Support & NOC Engineer. Stretch pack.
Current PATH line = Tier 3 IT Support & NOC Engineer.
ICT era = Network Operations Engineer. Intern off.
Run: python3 "stretch/Tier 3 IT Support NOC Engineer/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "Tier 3 IT Support and NOC Engineer  |  Networks, infrastructure, monitoring and RCA  |  24x7",
"metrics": [
    ("120+", "users / sites"),
    ("~98%", "uptime"),
    ("Tier 3", "escalation / RCA"),
    ("24x7", "on-call ready"),
],
"show_referees": False,
"compact_education": True,
"summary": ("Tier 3 IT Support and NOC Engineer for a global organisation's Kenya operation: the "
    "escalation point when Tier 1 and Tier 2 cannot close it. I monitor infrastructure, isolate "
    "whether the fault is routing, switching, identity, endpoint, or the WAN, then I fix it, write "
    "the RCA, and change the config so it does not return. SolarWinds, PRTG, Nagios, Zabbix, Cosmos, "
    "Grafana. Cisco, Meraki, Ruckus, Aruba, Fortinet, Cisco ASA, WatchGuard. BGP, OSPF, EIGRP, STP, "
    "SD-WAN, MPLS, wireless. Wireshark and tcpdump. Microsoft 365, Azure/Entra ID, Active Directory, "
    "Windows and macOS. PowerShell and RMM. SIEM/EDR. Diagrams, SOPs, vendors, field technicians. "
    "Available for 24x7 operations and on-call."),
"skills": [
    ("NOC / monitoring", "SolarWinds, PRTG, Nagios, Zabbix, Cosmos, Grafana; alert triage, capacity, incident dashboards"),
    ("Routing / switching", "BGP, OSPF, EIGRP, STP, VLANs, LAN/WAN, MPLS, SD-WAN, wireless controllers, load balancers"),
    ("Packets / services", "Wireshark, tcpdump, TCP/IP, DNS, DHCP, NAT, QoS, VoIP/3CX"),
    ("Platforms / firewalls", "Cisco, Meraki, Ruckus, Aruba/HP, Cisco ASA, Fortinet, WatchGuard, MikroTik, VPN"),
    ("Identity / endpoints", "Windows, macOS, Active Directory, Azure/Entra ID, Microsoft 365, MFA, patching, firmware"),
    ("Ops / security", "RCA, Tier 1/2 escalation, PowerShell, RMM, SIEM/EDR, SOPs, network diagrams, DR, vendors, field techs, PMS/hospitality site ops"),
],
"experience": [
    {"org": "PATH", "title": "Tier 3 IT Support & NOC Engineer", "dates": "Mar 2025 to Present",
     "loc": "Kenya (Hybrid / Remote)  ·  24x7 on-call", "bullets": [
        "Primary technical escalation for Tier 1 and Tier 2 across multi-site infrastructure. Own high-impact incidents through RCA, restore service, and leave the change and the diagram behind. Uptime above 98%.",
        "Monitor production on SolarWinds, PRTG, Nagios, Zabbix, Cosmos, and Grafana. Act on SIEM/EDR and security-log alerts instead of waiting for a user ticket.",
        "Diagnose LAN, WAN, MPLS, SD-WAN, and wireless. Routing and switching: BGP, OSPF, EIGRP, STP. Packet analysis with Wireshark and tcpdump. DNS, DHCP, NAT, VoIP, QoS.",
        "Configure and manage routers, switches, firewalls (Cisco ASA, Fortinet, WatchGuard, Meraki MX), VLANs, VPNs, wireless controllers, and load balancers. Firmware, patching, config change with a backout.",
        "Support Windows, macOS, Active Directory, Azure/Entra ID, and Microsoft 365. PowerShell and RMM for repeat work. Coordinate vendors and field technicians. SOPs, diagrams, incident reports, DR and capacity notes.",
    ]},
    {"org": "PATH", "title": "Network Operations Engineer", "dates": "Jan 2024 to Mar 2025",
     "loc": "Kenya (multi-site)", "bullets": [
        "Ran network operations for 120+ users across sites: Cisco Meraki, Fortinet, MikroTik, VLANs, VPN, wireless, DNS/DHCP. Ticket response about 25% faster once the queue had a real escalation path.",
        "Stood up identity and endpoints as part of the same NOC: Active Directory, Entra ID, Microsoft 365, Windows and macOS, MFA, imaging, patching.",
        "Took the incidents Tier 1 could not close: routing, firewall, WAN, VoIP/3CX. Wrote the RCA and the SOP. Trained 80+ users so noise dropped out of the queue.",
        "Automation with PowerShell for provisioning, patching, and reporting. Vendor and field-tech coordination when the break needed hands on site.",
    ]},
    {"org": "Independent", "title": "NOC tooling and automation (project-based)", "dates": "2023 to Present",
     "loc": "Remote", "bullets": [
        "Built an operations platform with 40+ REST APIs: incident/NOC tickets, IT asset lifecycle, dashboards, role-based access. Grafana-style visibility so the queue sees the estate, not a spreadsheet.",
        "Documented configuration, diagrams, and runbooks so the next engineer can take a shift without a handover call.",
    ]},
],
"certs": "CCNA  |  AZ-900  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the Tier 3 IT Support and NOC Engineer role. I am the escalation point. When Tier 1 and Tier 2 cannot close it, I take the incident, isolate the layer (routing, switching, WAN, identity, endpoint, or application), restore service, and write the RCA so the same fault does not come back.",
    "I monitor infrastructure on SolarWinds, PRTG, Nagios, Zabbix, Cosmos, and Grafana. I work Cisco, Meraki, Ruckus, Aruba, Fortinet, Cisco ASA, and WatchGuard. BGP, OSPF, EIGRP, STP, SD-WAN, MPLS, wireless. Wireshark and tcpdump. DNS, DHCP, NAT, VoIP, QoS. Windows, macOS, Active Directory, Azure/Entra ID, Microsoft 365. PowerShell and RMM. SIEM and EDR. Diagrams, SOPs, vendors, field technicians, hospitality-style site ops.",
    "At PATH I have kept multi-site infrastructure above 98% uptime for 120+ users, including 24x7-style on-call. I hold a BSc in Applied Computing and CCNA plus Azure Fundamentals (AZ-900). I am ready to take the on-call rota from week one.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]


def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "Tier 3 IT Support and NOC Engineer")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)


if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
