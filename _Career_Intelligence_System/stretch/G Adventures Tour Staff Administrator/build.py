#!/usr/bin/env python3
"""Tour Staff Administrator — G Adventures. Stretch pack.
ECD = Field Programme Administrator. ICT Associate era = Regional Operations and Staff Administrator.
HMIS Officer unchanged. Intern off.
Run: python3 "stretch/G Adventures Tour Staff Administrator/build.py"
"""
import os, sys
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESUMES = os.path.abspath(os.path.join(HERE, "..", "..", "resumes"))
sys.path.insert(0, RESUMES)
import build_resumes as br
import pdf_builder as pb

cv = {
"headline": "Tour Staff Administrator  |  Field staff records, logistics and systems  |  East Africa",
"metrics": [
    ("120+", "field staff supported"),
    ("~98%", "ops uptime"),
    ("Africa", "Kenya / Zambia"),
    ("On-call", "after hours"),
],
"show_referees": False,
"compact_education": True,
"summary": ("Administrator supporting field staff and regional operations in East Africa. Keeps people, "
    "files and systems current: staff records, onboarding, team channels, site logistics, documentation "
    "and after-hours operational cover. Concurrent PATH work: Field Programme Administrator on Nuru ya Mtoto "
    "(ECD) and Health Informatics / HMIS Officer. Microsoft 365, Google Workspace, Sheets, Teams, Freshdesk. "
    "English, used to tight operational deadlines and a flexible rota."),
"skills": [
    ("Staff records", "Onboarding and offboarding files, contact profiles, access and expiry-style tracking, digital filing"),
    ("Field ops", "Site logistics, vendor follow-up, training and meeting admin, multi-site East Africa including Zambia touchpoints"),
    ("Systems and comms", "Microsoft 365, Google Workspace, Sheets, Teams channels, SharePoint / Drive, Freshdesk, portals"),
    ("Pressure", "After-hours incident rotation, clear English, status before anyone has to chase"),
    ("Documentation", "SOPs, trip-style packs, knowledge base, receipts and transaction notes, monthly reporting"),
],
"experience": [
    {"org":"PATH","title":"Field Programme Administrator, Nuru ya Mtoto (ECD)","dates":"Mar 2025 to Present","loc":"Concurrent with HMIS Officer · Kenya (Hybrid / Remote)","bullets":[
        "Compile and keep field packs current for programme staff: trackers, contacts, notes and files so the next trip or forum does not start from zero.",
        "Follow up across field and office teams, including Zambia ECD learning touchpoints. Flag delays in writing and escalate to the lead.",
        "Support meetings and training: attendance, logistics and the write-up afterwards.",
    ]},
    {"org":"PATH","title":"Health Informatics / HMIS Officer","dates":"Mar 2025 to Present","loc":"Kisumu/Nyamira, Kenya (Hybrid / Remote)","bullets":[
        "Keep KenyaEMR/TaifaCare and DHIS2 running for field sites at above 98% uptime. Stay calm when something is down, isolate the layer, restore, document.",
        "Maintain operational databases, reporting cycles and HIS documentation. Reporting effort down about 40% through tighter Excel workflows.",
    ]},
    {"org":"PATH","title":"Regional Operations & Staff Administrator","dates":"Jan 2024 to Mar 2025","loc":"Kisumu, Kenya (multi-site)","bullets":[
        "Ran day-to-day staff administration for 120+ people across sites: records, onboarding, offboarding, Teams channels and a live task board (Freshdesk). Response about 25% faster.",
        "Kept digital files, asset notes and knowledge-base articles accurate. Tracked who was current and who had left so access matched the roster.",
        "Coordinated site logistics and vendor follow-up (travel between sites, connectivity, equipment). Trained 80+ staff. Covered after-hours incidents on rotation.",
    ]},
    {"org":"Independent","title":"Projects & Consulting (project-based)","dates":"2023 to Present","loc":"Remote","bullets":[
        "Remote admin and systems work: Google Workspace / Microsoft 365, documentation, onboarding and follow-up until the client could run the file.",
    ]},
],
"certs": "AZ-900  |  Google IT Support  |  IBM Technical Support  |  GitHub Professional  |  Microsoft Data Analysis",
}

COVER_BODY = [
    "Dear Hiring Team,",
    "",
    "I am applying for the Tour Staff Administrator role. I have not worked at G Adventures and I have not been a CEO or tour leader. I have spent the last three years in East Africa keeping field staff and regional operations organised: records, channels, logistics, documentation, and an after-hours phone when something broke.",
    "At PATH I currently support Nuru ya Mtoto as Field Programme Administrator (ECD): field packs, follow-up, meeting admin, including Zambia touchpoints. Alongside that I am Health Informatics / HMIS Officer, which is the systems and crisis layer. Before that I was Regional Operations and Staff Administrator for 120+ people across Kenyan sites: onboarding files, Teams channels, site logistics, training, and after-hours cover.",
    "I work in English every day. Microsoft 365, Google Workspace, Sheets, Teams. I can work a flexible rota including evening operational cover. I care about this region. I want to be the person who makes sure the field team has the file, the channel, and the logistics before they need to shout.",
    "Thank you.",
    "",
    "Sincerely,",
    "Meshack Ariri",
]

def build_cover_letter(path):
    doc = Document()
    br.set_base_style(doc)
    br.add_name(doc, "MESHACK ARIRI")
    br.add_title(doc, "Tour Staff Administrator")
    br.add_contact(doc, br.CONTACT)
    for para in COVER_BODY:
        if para == "":
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(para)
        r.font.size = Pt(10.5)
    doc.save(path)
    print("Saved:", path)

if __name__ == "__main__":
    files = br.pack_files(HERE)
    br.build(cv, files["cv_docx"])
    build_cover_letter(files["cover_docx"])
    cover_title = cv["headline"].split("|")[0].strip()
    pb.build_cv_pdf(cv, files["cv_pdf"], br.CONTACT)
    pb.build_cover_pdf(COVER_BODY, files["cover_pdf"], br.CONTACT, cover_title)
    print("Generated:", HERE)
